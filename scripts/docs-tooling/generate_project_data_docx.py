from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement as SharedOxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


OUTPUT = "SentinelCV_Project_Data_Inventory.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        el = tc_mar.find(qn(f"w:{m}"))
        if el is None:
            el = OxmlElement(f"w:{m}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_cell(cell, bold=False, size=9.5, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        p.alignment = align
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.bold = False
    r.font.color.rgb = RGBColor(0, 0, 0)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(10)
        r2 = p2.add_run(subtitle)
        r2.font.name = "Calibri"
        r2.font.size = Pt(11)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(85, 85, 85)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = RGBColor(47, 84, 150)
    r.font.size = Pt(16 if level == 1 else 13 if level == 2 else 12)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0, 0, 0)


def set_table_style(table):
    table.autofit = False
    table.style = "Table Grid"
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert(0, tblW)
    tblW.set(qn("w:w"), "9360")
    tblW.set(qn("w:type"), "dxa")


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)

add_title(doc, "SentinelCV Project Data Inventory", "Database inventory, thesis-ready project data section, and UI/API data map")

add_heading(doc, "1. Full Database / Data Inventory", 1)
add_para(doc, "The following table summarizes the active database and data surfaces present in the SentinelCV codebase. Primary operational tables are listed first, followed by supporting and experimental data domains.")

table1 = doc.add_table(rows=1, cols=3)
set_table_style(table1)
set_col_widths(table1, [1.65, 1.45, 3.4])
hdr = table1.rows[0].cells
hdr[0].text = "Table / Data Surface"
hdr[1].text = "Purpose"
hdr[2].text = "Key Data / Fields"
set_repeat_table_header(table1.rows[0])
for c in hdr:
    set_cell_shading(c, "E8EEF5")
    set_cell_margins(c)
    style_cell(c, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)

rows = [
    ("organizations", "Tenant boundary", "id, name, thresholds, retention, notification flags, settings, created_at"),
    ("users", "App users / staff", "id, organization_id, email, full_name, password_hash, role, is_active, created_at"),
    ("refresh_token_sessions", "Refresh token rotation", "hashed token, JTI, expiry, revocation, replacement hash, IP, user agent"),
    ("audit_logs", "Event trail", "action, entity type, entity id, details, timestamp, org/user linkage"),
    ("visitors", "Visitor profiles", "name, email, phone, description, notes, metadata, is_known, is_active, detection history"),
    ("face_data", "Face embeddings / references", "visitor_id, embedding, image_url, quality_score, face_angle, is_primary, created_at"),
    ("cameras", "Camera registry", "name, rtsp_url, location, is_active, status, last_seen"),
    ("camera_groups", "Camera grouping", "group metadata and organization linkage"),
    ("camera_sessions", "Live session tracking", "camera, user, status, timestamps, runtime metadata"),
    ("video_processing_jobs", "Uploaded video jobs", "job state, progress, file metadata, timestamps"),
    ("visitor_logs", "Operational detection logs", "visitor_id, camera_id, face_data_id, paths, confidence, status, identified, timestamp, track_id"),
    ("detection_logs", "Frame/event detections", "session linkage, visitor linkage, confidence, timing, metadata"),
    ("visitor_alerts", "Visitor alerts", "alert type, severity, visitor/log linkage, state"),
    ("edge_devices", "Edge registry", "device identity, endpoint, tokens, sync/export metadata, metrics"),
    ("edge_device_events", "Edge events", "device linkage, event type, payload, timestamp"),
    ("webhooks", "Outbound integrations", "target URL, filters, secret, status"),
    ("webhook_logs", "Webhook history", "delivery result, response code, retry state"),
    ("api_keys", "API access keys", "key metadata, hashed secret, scopes, status"),
    ("notifications", "In-app notifications", "user/org linkage, type, content, read state"),
    ("alert_config / alert_rules", "Alerting settings", "thresholds, channels, triggers, actions"),
    ("liveness_challenges", "Challenge sessions", "challenge type, state, timestamps, visitor/log linkage"),
    ("liveness_scores", "Liveness outcomes", "score, result, challenge/session linkage"),
    ("consent_records", "Consent tracking", "consent type, status, timestamps, visitor linkage"),
    ("gdpr_requests", "GDPR requests", "request type, status, timestamps, processing metadata"),
    ("data_retention_policies", "Retention policy", "retention rules, scope, expiration behavior"),
    ("rate_limit_rules", "Rate limiting", "target, threshold, window, enabled state"),
    ("carbon_metrics", "Sustainability metrics", "emissions/usage metrics, org linkage, timestamps"),
    ("bias_audit_records", "Fairness audits", "audit result, metrics, reviewer, timestamps"),
    ("training_jobs", "Training pipeline", "training state, parameters, timestamps"),
    ("hpo_jobs", "Hyperparameter optimization", "optimization state, parameters, results"),
    ("ensembles", "Model ensembles", "ensemble metadata, components, state"),
    ("active_learning_jobs", "Active learning", "queue state, samples, timestamps"),
    ("learning_samples", "Learning samples", "source, label, confidence, metadata"),
    ("recognition_feedback", "Feedback loop", "confirmed/rejected, correction data"),
    ("analytics_predictions", "Prediction outputs", "score, labels, timestamps"),
    ("data_quality_configs", "Data quality settings", "thresholds, validation rules"),
    ("data_quality_audit_jobs", "Data quality audits", "run metadata, state, timestamps"),
    ("quality_findings", "Data quality findings", "finding type, severity, affected records"),
    ("demographic_analysis", "Demographic analytics", "analysis metrics, timestamps"),
    ("quality_trend_analysis", "Quality trends", "time-series quality metrics"),
    ("synthetic_data_configs", "Synthetic data settings", "generation config, state"),
    ("synthesis_jobs", "Synthetic generation jobs", "job state, inputs, outputs"),
    ("generated_images", "Generated images", "image refs, metadata, timestamps"),
    ("quality_metrics_synthetic", "Synthetic quality metrics", "quality scores, metrics"),
    ("temporal_augmentation_configs", "Temporal augmentation settings", "augmentation config, state"),
    ("augmentation_jobs", "Augmentation jobs", "job state, parameters, timestamps"),
    ("generated_sequences", "Generated sequences", "sequence metadata, outputs"),
    ("expression_metrics", "Expression metrics", "labels, confidence, summary stats"),
    ("model_versions", "Model registry", "model metadata, versioning, deployment state"),
    ("ab_test_experiments", "A/B experiments", "config, variants, status"),
    ("ab_test_results", "A/B results", "metrics, winner, timestamps"),
    ("federated_rounds / federated_updates", "Federated learning", "round metadata, aggregation state, client updates"),
    ("gait_signatures / voice_prints", "Biometric modalities", "gait and voice features linked to visitors"),
    ("multimodal_embeddings / audio_features / text_bio_features", "Multimodal profile data", "combined embeddings and auxiliary features"),
    ("multimodal_fusion_configs", "Fusion settings", "fusion strategy, thresholds"),
    ("three_d_face_data / three_d_face_comparisons", "3D face data", "3D geometry, comparisons, match state"),
    ("sso_connections / sso_providers / sso_sessions", "SSO data", "provider config, connection metadata, sessions"),
    ("ldap_configs / ldap_sync_logs", "LDAP data", "server config, sync history"),
    ("multi_angle_configs", "Multi-angle recognition", "angle thresholds, supported angles, strategy"),
    ("behavior_events", "Behavior analytics", "behavior labels, linkage"),
    ("cross_camera_movement_summaries", "Cross-camera movement", "transition summaries, visitor linkage"),
    ("continuous_learning_signals", "Continuous learning", "signals, visitor linkage"),
    ("vision_analytics_events", "Vision analytics", "analytics outputs, timestamps"),
    ("future_enhancements", "Roadmap records", "planned feature entries, status"),
    ("password_reset_tokens", "Password reset flow", "token, expiry, user linkage"),
    ("user_sessions", "Session metadata", "session state, timestamps, device metadata"),
]

for a, b, c in rows:
    row = table1.add_row().cells
    row[0].text = a
    row[1].text = b
    row[2].text = c
    for cell in row:
        set_cell_margins(cell)
        style_cell(cell, size=9.1)

add_heading(doc, "2. Project Data Section", 1)
for para in [
    "The Automated Visitor Tracking and Identification System stores and processes multiple categories of data to support authentication, face recognition, visitor monitoring, reporting, and administrative control. The database is organized around a multi-tenant structure, where each organization acts as a logical boundary for users, visitors, cameras, logs, alerts, and related operational records.",
    "The core administrative data includes organizations, users, refresh-token sessions, and audit logs. Organizations store tenant-level configuration such as confidence thresholds, retention policies, and notification preferences. Users represent authenticated staff members with role-based access control, while refresh-token sessions support secure login persistence and token rotation. Audit logs preserve a trace of important actions for accountability and security review.",
    "The main operational data revolves around visitors, face data, cameras, and visitor logs. Visitor records store identity and profile information, including descriptive metadata, known or unknown status, activity state, and detection history. Face data stores face embeddings and reference images used for recognition. Camera records keep track of live video sources, locations, status, and last-seen activity. Visitor logs capture the runtime output of the system, including the detected visitor, linked camera, confidence score, tracking identifier, timestamp, and identification status.",
    "The platform also stores live-session and video-processing data to support real-time monitoring and media-based analysis. Camera session records track active monitoring sessions, while video-processing jobs represent queued or completed processing tasks for uploaded footage. Detection logs store frame-level or event-level detection results that can later be reviewed or assigned to a known visitor.",
    "In addition to the core tracking data, the system maintains operational and governance data such as alerts, notifications, webhooks, API keys, consent records, GDPR requests, retention policies, and liveness challenge scores. These records support system automation, compliance, and security. The database further includes advanced and experimental data domains such as analytics predictions, training jobs, model versions, synthetic data generation, multimodal embeddings, three-dimensional face data, and federated learning records. These tables extend the system for research, model improvement, and future feature expansion.",
    "Overall, the data model is designed to support secure authentication, real-time visitor identification, evidence-backed logging, and scalable expansion into advanced AI-based monitoring and analytics.",
]:
    add_para(doc, para)

add_heading(doc, "3. UI Page / API Route Data Map", 1)
add_para(doc, "The table below maps the main frontend pages and route families to the data they read and write. It is based on the current frontend service layer and the mounted backend route families.")

table2 = doc.add_table(rows=1, cols=3)
set_table_style(table2)
set_col_widths(table2, [1.8, 2.65, 2.05])
hdr = table2.rows[0].cells
hdr[0].text = "UI Page / Feature"
hdr[1].text = "Reads"
hdr[2].text = "Writes"
set_repeat_table_header(table2.rows[0])
for c in hdr:
    set_cell_shading(c, "E8EEF5")
    set_cell_margins(c)
    style_cell(c, bold=True, size=9.5)

map_rows = [
    ("/login", "auth state, org bootstrap state", "login, register, forgot-password, reset-password"),
    ("/", "dashboard stats, visitor logs, cameras, realtime websocket stream", "session refresh and websocket join only"),
    ("/visitors", "visitor list, pagination, filters", "create, update, delete visitors"),
    ("/visitors/[id]", "visitor detail, face data, related logs, accuracy settings", "update visitor, upload face data, set primary face, delete face data, upload media, bulk upload, threshold updates"),
    ("/logs", "visitor logs, unidentified logs, dashboard stats", "assign log, create visitor from log, confirm/reject log, export logs"),
    ("/cameras and live views", "camera list, stream frame, camera health", "create, update, delete camera, test camera, group assignment"),
    ("/camera live session", "camera session state, detection logs", "start session, process frame, end session, manual assign"),
    ("/users", "user list, current user, sessions", "create, update, delete user, change password, revoke sessions"),
    ("/settings", "organization profile, org stats, alert config, alert rules", "update org settings, update alert config/rules"),
    ("/alerts", "alert configuration, alert rules, trigger history", "create, update, delete, activate, deactivate rules"),
    ("/notifications", "notification list and preferences", "update preferences, mark read/clear"),
    ("/audit-logs", "audit trail, login history", "export action only"),
    ("/reports", "reports data, filters, export output", "generate/export report requests"),
    ("/edge-devices", "device list, device details, telemetry/events", "create, update, delete device, rotate token, sync/export config"),
    ("/liveness", "liveness config, statistics, challenge history", "start challenge, verify challenge, update config"),
    ("/analytics", "analytics summaries, trends, flow metrics", "analytics generation triggers when used"),
    ("/models", "training stats, quality info, recommendations", "retrain, import directory, batch upload, quality check, augment"),
    ("/recommendations", "recommendation outputs, threshold suggestions", "none"),
    ("/accuracy", "accuracy stats, liveness config, anti-spoofing config", "update visitor threshold, update configs"),
    ("/compliance / GDPR", "retention settings, consent-related data, GDPR requests", "retention updates, export/delete requests, consent actions"),
    ("/settings/sso", "SSO provider config and sessions", "create, update, delete provider, auth actions"),
    ("/settings/ldap", "LDAP config and sync logs", "update config, test connection, sync"),
    ("/health / readiness / metrics / websocket", "runtime health, dependency readiness, metrics, realtime dashboard stream", "none"),
]

for a, b, c in map_rows:
    row = table2.add_row().cells
    row[0].text = a
    row[1].text = b
    row[2].text = c
    for cell in row:
        set_cell_margins(cell)
        style_cell(cell, size=9.1)

doc.save(OUTPUT)
print(OUTPUT)
