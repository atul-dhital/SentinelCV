from __future__ import annotations

import re
import sys
from pathlib import Path
from docx import Document


path = Path(sys.argv[1])
doc = Document(path)
headings = []
for i, p in enumerate(doc.paragraphs):
    m = re.fullmatch(r"Heading ([1-9])", p.style.name)
    if m:
        headings.append((i, int(m.group(1)), " ".join(p.text.split())))

print("HEADING TREE")
stack = []
previous_level = 0
for i, level, text in headings:
    issue = ""
    if previous_level and level > previous_level + 1:
        issue = " [SKIPPED LEVEL]"
    while stack and stack[-1][0] >= level:
        stack.pop()
    parent = stack[-1][2] if stack else "<root>"
    print(f"{i:04d} H{level} parent={parent!r} text={text!r}{issue}")
    stack.append((level, i, text))
    previous_level = level

print("\nPOTENTIAL FAKE HEADINGS")
heading_texts = {text for _, _, text in headings}
for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if not text or text in heading_texts or p.style.name in {"Caption", "table of figures"} or p.style.name.startswith("TOC"):
        continue
    if len(text) > 90 or text.endswith((".", ":", ";", "?", "!")):
        continue
    textual_runs = [r for r in p.runs if r.text.strip()]
    if not textual_runs:
        continue
    all_bold = all(bool(r.bold) for r in textual_runs)
    all_italic = all(bool(r.italic) for r in textual_runs)
    if all_bold or all_italic:
        print(f"{i:04d} style={p.style.name!r} bold={all_bold} italic={all_italic} text={text!r}")
