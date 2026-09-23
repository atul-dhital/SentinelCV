from __future__ import annotations

import hashlib
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SOURCE = Path(r"D:\college\Final-Year-Documentation\docs\final.docx")
OUTPUT = Path(r"D:\college\Final-Year-Documentation\docs\final_APA7.docx")

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
BLACK = RGBColor(0, 0, 0)


def body_text_signature(doc: Document) -> tuple[int, str]:
    paragraphs = []
    for p_el in doc.element.body.xpath(".//w:p"):
        paragraphs.append("".join(p_el.xpath(".//w:t/text()")))
    joined = " ".join(paragraphs)
    joined = re.sub(r"(?i)(?<!\w)(?:ii|iii|iv)\.\s+", " ", joined)
    joined = re.sub(r"\s+", " ", joined).strip()
    return len(joined), hashlib.sha256(joined.encode("utf-8")).hexdigest()


def set_font(font, *, bold=None, italic=None):
    font.name = FONT_NAME
    font.size = FONT_SIZE
    font.color.rgb = BLACK
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def set_style_font(style, *, bold=None, italic=None):
    set_font(style.font, bold=bold, italic=italic)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_NAME)


def set_run_font(run_element):
    rpr = run_element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_NAME)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), "24")


def set_double_spacing(paragraph, *, first_line=None, left=None, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    pf = paragraph.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = alignment
    if first_line is not None:
        pf.first_line_indent = first_line
    if left is not None:
        pf.left_indent = left


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, cached, fld_end):
        run._r.append(el)
    set_run_font(run._r)


def reset_container(container):
    # Remove all legacy header/footer content, including fields hidden inside
    # text boxes that python-docx does not expose through .paragraphs.
    for child in list(container._element):
        container._element.remove(child)
    return container.add_paragraph()


def configure_styles(doc: Document):
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE):
            try:
                set_style_font(style)
            except (AttributeError, ValueError):
                pass

    normal = doc.styles["Normal"]
    set_style_font(normal)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    if "Title" in doc.styles:
        title = doc.styles["Title"]
        set_style_font(title, bold=True, italic=False)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.line_spacing = 2.0
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(0)
        title.paragraph_format.first_line_indent = Inches(0)

    heading_specs = {
        "Heading 1": (WD_ALIGN_PARAGRAPH.CENTER, Inches(0), True, False),
        "Heading 2": (WD_ALIGN_PARAGRAPH.LEFT, Inches(0), True, False),
        "Heading 3": (WD_ALIGN_PARAGRAPH.LEFT, Inches(0), True, True),
        "Heading 4": (WD_ALIGN_PARAGRAPH.LEFT, Inches(0.5), True, False),
        "Heading 5": (WD_ALIGN_PARAGRAPH.LEFT, Inches(0.5), True, True),
    }
    for name, (align, indent, bold, italic) in heading_specs.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        set_style_font(style, bold=bold, italic=italic)
        pf = style.paragraph_format
        pf.alignment = align
        pf.left_indent = indent
        pf.first_line_indent = Inches(0)
        pf.line_spacing = 2.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.keep_with_next = True
        if name == "Heading 1":
            pf.page_break_before = True

    for name in ("Caption", "table of figures", "TOC 1", "TOC 2", "TOC 3"):
        if name in doc.styles:
            style = doc.styles[name]
            set_style_font(style)
            style.paragraph_format.line_spacing = 2.0
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.first_line_indent = Inches(0)


def configure_sections(doc: Document):
    doc.settings.odd_and_even_pages_header_footer = False
    for idx, section in enumerate(doc.sections):
        landscape = section.orientation == WD_ORIENT.LANDSCAPE or section.page_width > section.page_height
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if landscape:
            section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        else:
            section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    for idx, section in enumerate(doc.sections):
        header_p = reset_container(section.header)
        set_double_spacing(header_p, first_line=Inches(0), alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        add_page_field(header_p)
        footer_p = reset_container(section.footer)
        set_double_spacing(footer_p, first_line=Inches(0), alignment=WD_ALIGN_PARAGRAPH.LEFT)

        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if idx in (0, 1):
            if pg_num is None:
                pg_num = OxmlElement("w:pgNumType")
                sect_pr.append(pg_num)
            pg_num.set(qn("w:start"), "1")
        elif pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num)
        if idx > 1 and qn("w:start") in pg_num.attrib:
            del pg_num.attrib[qn("w:start")]
        pg_num.set(qn("w:fmt"), "lowerRoman" if idx == 0 else "decimal")


def configure_settings(doc: Document):
    settings = doc.settings._element
    for tag, val in (("w:autoHyphenation", "0"), ("w:doNotHyphenateCaps", "1")):
        el = settings.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            settings.append(el)
        el.set(qn("w:val"), val)
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def format_paragraphs(doc: Document):
    title_page_last = 13
    abstract_body_index = 21

    for idx, paragraph in enumerate(doc.paragraphs):
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()

        if idx <= title_page_last:
            set_double_spacing(paragraph, first_line=Inches(0), alignment=WD_ALIGN_PARAGRAPH.CENTER)
            if idx in (2, 3):
                for run in paragraph.runs:
                    run.bold = True
            continue

        if style_name == "Heading 1":
            set_double_spacing(paragraph, first_line=Inches(0), left=Inches(0), alignment=WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name == "Heading 2":
            set_double_spacing(paragraph, first_line=Inches(0), left=Inches(0), alignment=WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name == "Heading 3":
            set_double_spacing(paragraph, first_line=Inches(0), left=Inches(0), alignment=WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name == "Heading 4":
            set_double_spacing(paragraph, first_line=Inches(0), left=Inches(0.5), alignment=WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name == "Heading 5":
            set_double_spacing(paragraph, first_line=Inches(0), left=Inches(0.5), alignment=WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.keep_with_next = True
        elif style_name in ("Caption", "table of figures") or style_name.startswith("TOC"):
            set_double_spacing(paragraph, first_line=Inches(0), alignment=WD_ALIGN_PARAGRAPH.LEFT)
        elif style_name in ("List Paragraph", "List Bullet", "List Number"):
            level = 0
            if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                ilvl = paragraph._p.pPr.numPr.ilvl
                if ilvl is not None:
                    level = int(ilvl.val)
            set_double_spacing(
                paragraph,
                first_line=Inches(-0.25),
                left=Inches(0.5 + 0.25 * level),
                alignment=WD_ALIGN_PARAGRAPH.LEFT,
            )
        else:
            set_double_spacing(paragraph, first_line=Inches(0.5), alignment=WD_ALIGN_PARAGRAPH.LEFT)

        if idx == abstract_body_index:
            paragraph.paragraph_format.first_line_indent = Inches(0)

        if re.match(r"^(Figure|Table)\s+\d+\b", text, flags=re.IGNORECASE):
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def format_all_runs(doc: Document):
    for run_el in doc.element.xpath(".//w:r"):
        set_run_font(run_el)


def apply_reference_hanging_indents(doc: Document):
    in_references = False
    for p_el in doc.element.body.xpath(".//w:p"):
        text = "".join(p_el.xpath(".//w:t/text()")).strip()
        style_vals = p_el.xpath("./w:pPr/w:pStyle/@w:val")
        style = style_vals[0] if style_vals else ""
        normalized = re.sub(r"\s+", " ", text).strip().lower()
        if normalized in {"references", "bibliography"}:
            in_references = True
            continue
        if normalized in {"appendix", "appendices"} or normalized.startswith("appendix a"):
            in_references = False
        if not in_references or not text or style.lower().startswith("heading"):
            continue
        ppr = p_el.get_or_add_pPr()
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "720")
        ind.attrib.pop(qn("w:firstLine"), None)


def format_caption_runs(doc: Document):
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "Caption":
            continue
        text = paragraph.text
        match = re.match(r"^((?:Figure|Table)\s+\d+)(.*)$", text, flags=re.IGNORECASE | re.DOTALL)
        if not match or any(run._r.xpath(".//w:fldChar|.//w:drawing|.//w:pict") for run in paragraph.runs):
            continue
        prefix, rest = match.groups()
        for run in paragraph.runs:
            run.text = ""
        first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        first.text = prefix
        first.bold = True
        first.italic = False
        if rest:
            second = paragraph.add_run(rest)
            second.bold = False
            second.italic = True


def set_numbering_definition(doc: Document, num_id: str, num_format: str, level_text: str, ilvl: str = "0"):
    root = doc.part.numbering_part.element
    nums = root.xpath(f'./w:num[@w:numId="{num_id}"]')
    if not nums:
        raise RuntimeError(f"Numbering definition {num_id} was not found")
    abstract_id = nums[0].xpath('./w:abstractNumId/@w:val')[0]
    abstracts = root.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]')
    lvl = next(el for el in abstracts[0].findall(qn("w:lvl")) if el.get(qn("w:ilvl")) == ilvl)
    num_fmt = lvl.find(qn("w:numFmt"))
    lvl_text = lvl.find(qn("w:lvlText"))
    num_fmt.set(qn("w:val"), num_format)
    lvl_text.set(qn("w:val"), level_text)
    suff = lvl.find(qn("w:suff"))
    if suff is None:
        suff = OxmlElement("w:suff")
        lvl.append(suff)
    suff.set(qn("w:val"), "space")
    ppr = lvl.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        lvl.append(ppr)
    tabs = ppr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        ppr.append(tabs)
    for child in list(tabs):
        tabs.remove(child)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    if num_format == "bullet":
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), FONT_NAME)
        rfonts.set(qn("w:hAnsi"), FONT_NAME)


def replace_paragraph_text(paragraph, text: str):
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    label, sep, remainder = text.partition(":")
    if sep:
        run.text = label + sep
        run.bold = True
        paragraph.add_run(remainder)
    else:
        run.text = text


def insert_list_paragraph_after(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_p.append(deepcopy(paragraph._p.pPr))
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    replace_paragraph_text(new_paragraph, text)
    return new_paragraph


def normalize_list_types(doc: Document):
    # Unordered characteristics and Agile benefits become bullets.
    set_numbering_definition(doc, "2", "bullet", "\u2022")
    set_numbering_definition(doc, "30", "bullet", "\u2022")
    # Ordinary research-method and findings lists are unordered; APA uses bullets,
    # not Roman numerals, for these groups.
    set_numbering_definition(doc, "42", "bullet", "\u2022", ilvl="2")
    for num_id in ("44", "45", "46"):
        set_numbering_definition(doc, num_id, "bullet", "\u2022")
    # The end-to-end AI pipeline is a true sequence, so use Arabic numbers.
    set_numbering_definition(doc, "10", "decimal", "%1.")

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("ii. Staff User / Security Personnel:"):
            for run in paragraph.runs:
                if "ii. Staff User / Security Personnel:" in run.text:
                    run.text = run.text.replace("ii. Staff User / Security Personnel:", "Staff User / Security Personnel:", 1)
                    break

    agile = next(
        (p for p in doc.paragraphs if p.text.startswith("Incremental Development:") and "ii. Flexibility in Model Tuning:" in p.text),
        None,
    )
    if agile is not None:
        parts = re.split(r"\n(?:ii|iii|iv)\.\s+", agile.text)
        replace_paragraph_text(agile, parts[0])
        cursor = agile
        for part in parts[1:]:
            cursor = insert_list_paragraph_after(cursor, part)


def suppress_heading_number(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)
    numpr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    numpr.append(num_id)
    ppr.append(numpr)


def clear_direct_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def normalize_heading_hierarchy(doc: Document):
    title = next(
        (p for p in doc.paragraphs if p.text.strip() == "Automated Visitor Tracking and Identification System"),
        None,
    )
    if title is not None:
        title.style = doc.styles["Title"]

    numbered_h3 = {
        "Organization Administrators.",
        "For Staff Users / Security Personnel",
        "For the System (Technical Deliverables)",
        "Justification for Python",
        "Justification",
    }
    numbered_h4 = {
        "Hardware",
        "Software",
        "Software Stack and Tools",
        "Planning Phase",
        "Requirement Analysis Phase",
        "Design Phase",
        "Implementation Phase",
        "Testing Phase",
        "Deployment Phase",
        "Evaluation & Feedback Phase",
        "Performance Benchmarking:",
        "Structured Surveys",
        "Functional Requirements (FR):",
        "Non-Functional Requirements (NFR):",
        "User Interface and Experience:",
    }
    numbered_h5 = {
        "Testing Procedure:",
        "Expectations Drawn From the Survey and Interviews:",
        "System Enhancements:",
        "Outcome:",
    }
    unnumbered_h3 = {
        "a) System Metrics Collection",
        "b) Surveys",
        "c) Interviews and Observational Studies",
    }
    unnumbered_h4 = {
        "Video Processing Benchmarks:",
        "Detection and Recognition Accuracy:",
    }

    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        target_style = None
        suppress = False
        if text in numbered_h3:
            target_style = "Heading 3"
        elif text in numbered_h4:
            target_style = "Heading 4"
        elif text in numbered_h5:
            target_style = "Heading 5"
        elif text in unnumbered_h3:
            target_style = "Heading 3"
            suppress = True
        elif text in unnumbered_h4:
            target_style = "Heading 4"
            suppress = True
        if target_style is None:
            continue
        paragraph.style = doc.styles[target_style]
        clear_direct_numbering(paragraph)
        if suppress:
            suppress_heading_number(paragraph)


def main():
    doc = Document(SOURCE)
    before = body_text_signature(doc)
    configure_styles(doc)
    configure_sections(doc)
    configure_settings(doc)
    normalize_list_types(doc)
    normalize_heading_hierarchy(doc)
    format_paragraphs(doc)
    format_caption_runs(doc)
    apply_reference_hanging_indents(doc)
    format_all_runs(doc)
    doc.save(OUTPUT)

    check = Document(OUTPUT)
    after = body_text_signature(check)
    if before != after:
        raise RuntimeError(f"Body text changed: before={before}, after={after}")
    print(f"saved={OUTPUT}")
    print(f"body_text_nodes={before[0]} sha256={before[1]}")
    print(f"sections={len(check.sections)} paragraphs={len(check.paragraphs)} tables={len(check.tables)}")


if __name__ == "__main__":
    main()
