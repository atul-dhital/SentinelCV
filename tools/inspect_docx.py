from collections import Counter
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"D:\college\Final-Year-Documentation\docs\final.docx")
doc = Document(SOURCE)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print("styles:", Counter(p.style.name for p in doc.paragraphs).most_common(20))

for i, section in enumerate(doc.sections):
    print(
        f"section {i}: size={section.page_width.inches:.2f}x{section.page_height.inches:.2f} "
        f"margins={section.top_margin.inches:.2f},{section.right_margin.inches:.2f},"
        f"{section.bottom_margin.inches:.2f},{section.left_margin.inches:.2f} "
        f"header={section.header_distance.inches:.2f} footer={section.footer_distance.inches:.2f}"
    )
    print(" header:", " | ".join(p.text for p in section.header.paragraphs if p.text.strip())[:500])
    print(" footer:", " | ".join(p.text for p in section.footer.paragraphs if p.text.strip())[:500])

print("FIRST 160 NONEMPTY PARAGRAPHS")
shown = 0
for idx, p in enumerate(doc.paragraphs):
    txt = " ".join(p.text.split())
    if not txt:
        continue
    print(f"{idx:04d} [{p.style.name}] {txt[:220]}")
    shown += 1
    if shown >= 160:
        break

print("POSSIBLE STRUCTURAL HEADINGS")
for idx, p in enumerate(doc.paragraphs):
    txt = " ".join(p.text.split())
    if txt and (p.style.name.lower().startswith("heading") or txt.lower() in {
        "abstract", "references", "bibliography", "appendix", "appendices",
        "table of contents", "list of figures", "list of tables"
    }):
        print(f"{idx:04d} [{p.style.name}] {txt[:250]}")

settings = doc.settings._element
print("auto_hyphenation_present=", settings.find(qn("w:autoHyphenation")) is not None)
